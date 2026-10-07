"""
Turns a file loaded from Studie+ into content Claude can read: text, extracted
PDF/DOCX text, or an image. Studie+ serves every file as application/octet-stream,
so the type is guessed from the file name.
"""
import base64
import io
import mimetypes

from fastmcp.tools import ToolResult
from fastmcp.utilities.types import Image

MAX_TEXT_CHARS = 100_000
MAX_IMAGE_BYTES = 5 * 1024 * 1024
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
GENERIC_TYPES = ("", "application/octet-stream", "binary/octet-stream")
TEXT_TYPES = ("application/json", "application/xml", "application/javascript")
IMAGE_TYPES = ("image/png", "image/jpeg", "image/gif", "image/webp")


def readable_file(result: dict) -> ToolResult:
    """Convert a load_lesson_file result into a tool result Claude can read."""
    info = {k: v for k, v in result.items() if k != "content"}
    if not result.get("success"):
        return ToolResult(content=info)

    content_type = file_type(result.get("file_name", ""), result.get("content_type", ""))
    info["content_type"] = content_type

    if result.get("is_text"):
        return _text_result(info, result["content"])

    data = base64.b64decode(result["content"])

    if content_type.startswith("text/") or content_type in TEXT_TYPES:
        return _text_result(info, data.decode("utf-8", errors="replace"))

    if content_type in IMAGE_TYPES:
        if len(data) > MAX_IMAGE_BYTES:
            return _unreadable(info, f"Billedet er større end {MAX_IMAGE_BYTES // (1024 * 1024)} MB.")
        return ToolResult(content=[info, Image(data=data, format=content_type.split("/")[1])])

    if content_type == "application/pdf":
        try:
            pages = _pdf_pages(data)
        except Exception as e:
            return _unreadable(info, f"PDF'en kunne ikke læses: {e}")
        if not any(page.strip() for page in pages):
            return _unreadable(info, "PDF'en indeholder ingen tekst, måske er den scannet.")
        return _text_result(info, "\n\n".join(f"--- Side {i} ---\n{text}" for i, text in enumerate(pages, 1)))

    if content_type == DOCX_TYPE:
        try:
            return _text_result(info, _docx_text(data))
        except Exception as e:
            return _unreadable(info, f"Word-dokumentet kunne ikke læses: {e}")

    return _unreadable(info, f"Filtypen {content_type} kan ikke vises direkte.")


def file_type(file_name: str, content_type: str) -> str:
    content_type = content_type.split(";")[0].strip().lower()
    if content_type not in GENERIC_TYPES:
        return content_type
    guessed, _ = mimetypes.guess_type(file_name)
    return guessed or "application/octet-stream"


def _text_result(info: dict, text: str) -> ToolResult:
    if len(text) > MAX_TEXT_CHARS:
        info["truncated"] = True
        info["total_chars"] = len(text)
        text = text[:MAX_TEXT_CHARS]
    return ToolResult(content={**info, "content": text})


def _unreadable(info: dict, reason: str) -> ToolResult:
    return ToolResult(content={
        **info,
        "success": False,
        "error": f"{reason} Brug download_lesson_file for at gemme filen på computeren.",
    })


def _pdf_pages(data: bytes) -> list:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [page.extract_text() or "" for page in reader.pages]


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
